# -*- coding: utf-8 -*-
"""live_report_probe.py —— 《店铺经营月报》成品 docx 的**实机**探针
==========================================================================
静态门禁 `verify_report_form.py` 验的是 md / HTML 这两层**源**；本探针验的是
用户真正拿到手、真正会传回来的那些 **docx**。两者不可互相替代：

  - 静态断言挡回退（改坏了立刻红）；
  - 实机探针挡"其实没生效"（HTML 写对了，但转换器把空格吃掉/把列表编上号）。

它**不进离线门禁**（离线门禁必须零外部依赖、可重复）：本探针要 python-docx、
要已生成好的 docx，还要两套解释器 —— 按本项目纪律属于"实机组"。

## 要验的三个 docx
    空表     C:\\Users\\ASUS\\Desktop\\店铺经营月报.docx           十项必须全 None
    填好     上面这份被 python-docx 就地填上数（模拟用户填表）      十项必须全中
    样例     C:\\Users\\ASUS\\Desktop\\店铺经营月报-填写样例.docx    十项必须全中
             —— 这份是**交付物本身**：它印着数字，所以它最可能"自己把自己的数字抢走"

## 为什么要两套解释器
    · python-docx 装在 3.12 / 托管 venv（`.pylibs` 是 cp314 轮子，3.12 import 不了）
    · 抽取器 `extract_*` / `_parse_*` 依赖 cp314 的 numpy 等
    · 所以：venv 负责"读 docx 出文本"，cp314 负责"跑真实抽取器"
      —— 中间用落盘 txt 交接（本项目反复踩到"stdout 被吞"，落盘最稳）

## 复刻的抽取路径
    `src/ui/app_chainlit.py :: _extract_docx` + `_norm_attach_text`
    先全部 `paragraphs`，再全部 `tables`（表格行按 ' | ' 拼）—— 注意
    `Document.paragraphs` **不含**表格单元格里的段落，所以视觉顺序 ≠ 抽取顺序。
    归一化把 [ \\t\\u00a0\\u3000]+ 压成一个空格。

## 最后一公里
    抽取通过后，用**抽出来的那十个数**跑一遍 `engine.store_diagnosis.diagnose_existing_store`
    —— 这一步回答的是"这份 docx 到底喂不喂得动经营测算专家"，而不只是"字段读没读到"。

## 用法（一条命令）
    "$HOME/.venv-html-to-docx/Scripts/python.exe" src/analysis/live_report_probe.py
    （会自动调用 cp314 的 --check 完成断言，rc=0 即通过）

可覆盖路径：
    ZL_REPORT_DOCX         默认 C:\\Users\\ASUS\\Desktop\\店铺经营月报.docx
    ZL_REPORT_SAMPLE_DOCX  默认 C:\\Users\\ASUS\\Desktop\\店铺经营月报-填写样例.docx
"""
import os
import re
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
sys.path.insert(0, str(HERE))

from report_sample import SAMPLE          # noqa: E402  （样例数值的唯一来源）

DOCX_EMPTY = Path(os.environ.get('ZL_REPORT_DOCX',
                                 r'C:\Users\ASUS\Desktop\店铺经营月报.docx'))
DOCX_SAMPLE = Path(os.environ.get('ZL_REPORT_SAMPLE_DOCX',
                                  r'C:\Users\ASUS\Desktop\店铺经营月报-填写样例.docx'))
DUMP_EMPTY = HERE / '_docx_dump_empty.txt'
DUMP_FILLED = HERE / '_docx_dump_filled.txt'
DUMP_SAMPLE = HERE / '_docx_dump_sample.txt'
PY314 = r'C:\Python314\python.exe'

# 模拟"用户填表"用的那组数 —— 与交付样例同一组（report_sample.SAMPLE）
LAB = [(lb, val) for lb, val, _ in SAMPLE]
EXPECT = {lb: exp for lb, _, exp in SAMPLE}


# ------------------------------------------------------------------ 1. 抽 docx
def dump(d):
    """逐字复刻 `_extract_docx`：段落在前、表格行在后。"""
    parts = [p.text for p in d.paragraphs]
    for tb in d.tables:
        for row in tb.rows:
            parts.append(' | '.join(c.text for c in row.cells))
    return '\n'.join(parts)


def norm(t):
    """逐字复刻 `_norm_attach_text`。"""
    t = re.sub(r'[ \t\u00a0\u3000]+', ' ', t or '')
    t = re.sub(r'\n\s*\n\s*\n+', '\n\n', t)
    return t.strip()


def rewrite(p, new_text):
    """整体替换段落文本（保留首个 run 的格式，清空其余）。"""
    if not p.runs:
        return False
    p.runs[0].text = new_text
    for r in p.runs[1:]:
        r.text = ''
    return True


def build_dumps():
    from docx import Document
    if not DOCX_EMPTY.exists():
        print(f'找不到空表 docx：{DOCX_EMPTY}')
        print('（先用 html-to-docx 把 docs/店铺经营月报-表单.html 转出来，'
              '或用 ZL_REPORT_DOCX 指定路径）')
        return 2

    d = Document(DOCX_EMPTY)
    DUMP_EMPTY.write_text(norm(dump(d)), encoding='utf-8')
    print(f'空表 dump ok：paragraphs={len(d.paragraphs)} tables={len(d.tables)}')

    # 造"用户填好"的副本：值写在标签后（顶格写）
    d2 = Document(DOCX_EMPTY)
    n_field = n_sum = 0
    for p in d2.paragraphs:
        t = p.text
        if t.strip().startswith('；月流水') or ('；月流水' in t and '\u3000' in t):
            it = iter([v for _, v in LAB])
            if rewrite(p, re.sub(r'\u3000+', lambda m: next(it, m.group()), t)):
                n_sum += 1
            continue
        for lab, val in LAB:
            if not t.strip().startswith(lab + '：'):
                continue
            m = re.search(r'\u3000+', t)
            if m and rewrite(p, t[:m.start()] + val + t[m.end():]):
                n_field += 1
            break
    print(f'填好 dump ok：字段行={n_field} 摘要行={n_sum}')
    DUMP_FILLED.write_text(norm(dump(d2)), encoding='utf-8')

    # 交付样例：已由 make_sample_report.py 填好，只读不造
    if not DOCX_SAMPLE.exists():
        print(f'⚠️ 找不到样例 docx：{DOCX_SAMPLE}（跳过样例那一组）')
        DUMP_SAMPLE.unlink(missing_ok=True)
    else:
        ds = Document(DOCX_SAMPLE)
        DUMP_SAMPLE.write_text(norm(dump(ds)), encoding='utf-8')
        print(f'样例 dump ok：paragraphs={len(ds.paragraphs)} tables={len(ds.tables)}')
    return 0


# ------------------------------------------------------------------ 2. 跑抽取器
def check():
    sys.path.insert(0, str(ROOT / 'src'))
    sys.path.insert(0, str(ROOT / '.pylibs'))
    from agent.agent import (                                  # noqa: E402
        extract_category, extract_rent, extract_monthly_revenue,
        extract_area, extract_brand)
    from agent.agent_graph import (                             # noqa: E402
        extract_city, _parse_staff, _parse_investment,
        _parse_price, _parse_daily_orders)

    fields = [('品类', extract_category), ('月流水', extract_monthly_revenue),
              ('月租金', extract_rent), ('经营面积', extract_area),
              ('店铺所在城市', extract_city), ('加盟品牌', extract_brand),
              ('全职员工', _parse_staff), ('前期投入', None),
              ('每单金额', _parse_price), ('日单量', _parse_daily_orders)]

    def run(text):
        got = {lb: fn(text) for lb, fn in fields if fn}
        got['前期投入'] = _parse_investment(text, strict=True)
        return got

    fails = []
    cases = [('空表', DUMP_EMPTY, None),
             ('填好', DUMP_FILLED, EXPECT),
             ('样例', DUMP_SAMPLE, EXPECT)]
    got_sample = None
    for tag, path, expect in cases:
        if not path.exists():
            if tag == '样例':
                print(f'  SKIP  {tag}：没有 {path.name}（桌面上没有样例 docx？）')
                continue
            print(f'  FAIL  {tag}：找不到 {path.name}')
            fails.append(f'{tag}:缺 dump')
            continue
        text = path.read_text(encoding='utf-8')
        print(f'== {tag}（{len(text)} 字符）')
        got = run(text)
        if tag == '样例':
            got_sample = got
        for lb, _ in fields:
            if expect is None:
                ok, want = got[lb] is None, 'None'
            else:
                ok, want = got[lb] == expect[lb], repr(expect[lb])
            print(f'  {"PASS" if ok else "FAIL"}  {tag}·{lb} 应为 {want}'
                  f'  |  实得 {got[lb]!r}')
            if not ok:
                fails.append(f'{tag}·{lb}')

    # ---- 最后一公里：抽出来的数直接喂引擎，看能不能跑出诊断 ----
    if got_sample:
        try:
            sys.path.insert(0, str(ROOT / 'src'))
            from engine.store_diagnosis import diagnose_existing_store  # noqa: E402
            r = diagnose_existing_store(
                category=got_sample['品类'],
                monthly_revenue=got_sample['月流水'],
                monthly_rent=got_sample['月租金'],
                area_m2=got_sample['经营面积'],
                city=got_sample['店铺所在城市'],
                staff=got_sample['全职员工'],
                investment=got_sample['前期投入'],
                brand=got_sample['加盟品牌'],
                price=got_sample['每单金额'],
                daily_orders=got_sample['日单量'])
            print('== 端到端：样例 docx → 抽取器 → 经营诊断')
            print(f'  月成本合计 {r["成本拆解"]["月成本合计"]}  月净利 {r["月净利"]}'
                  f'  净利率 {r["净利率"]}')
            print(f'  安全边际率 {r["安全边际率"]}  租金余量 {r["租金余量"]}')
            print(f'  品牌对标 {r["品牌对标"] and r["品牌对标"]["你/基准"]}')
            ok = r['月净利'] is not None and r['月净利'] > 0
            print(f'  {"PASS" if ok else "FAIL"}  样例是一份"正常经营"的店'
                  f'（月净利 > 0）  |  实得 {r["月净利"]!r}')
            if not ok:
                fails.append('样例·月净利')
        except Exception as e:                    # 引擎真挂了也得报出来，不能吞
            print(f'  FAIL  端到端调用引擎抛异常：{type(e).__name__}: {e}')
            fails.append('样例·端到端')

    print(f'FAILED={len(fails)}' + (f'  {fails}' if fails else ''))
    return 1 if fails else 0


if __name__ == '__main__':
    if '--check' in sys.argv:
        sys.exit(check())
    rc = build_dumps()
    if rc:
        sys.exit(rc)
    print(f'--- 交给 {PY314} 跑真实抽取器 ---')
    sys.exit(subprocess.run([PY314, str(Path(__file__).resolve()), '--check'],
                            cwd=str(ROOT)).returncode)
